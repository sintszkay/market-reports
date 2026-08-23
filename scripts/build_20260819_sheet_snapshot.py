#!/usr/bin/env python3
"""Build a point-in-time Google-Sheet-compatible snapshot for the missing 8/19 report."""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(r"C:\Users\sints\Documents\Codex\2026-08-11\us\work\market_watch_verify.xlsx")


def ticker_from_label(value):
    text = str(value or "")
    match = re.search(r"\(([A-Z0-9.]+)\)\s*$", text) or re.match(r"^([A-Z0-9.]+)\s*:", text)
    return match.group(1) if match else None


def percent(value):
    return f"{value:+.1f}%" if value is not None else None


def status(row):
    if row["above20"] and row["above50"] and row["above200"]:
        return "🟢 Strong Trend"
    if row["above50"] and row["above200"]:
        return "🟢 Uptrend"
    return "🔴 Downtrend"


def sheet_rows(rows, labels, universe):
    header = ["Sector", "Latest Price", "1D", None, "1M", None, "3M", None, "YTD", None,
              "Off % 52W High", None, "RSI", "Trend 20·50·200", "Status"]
    output = [header]
    by_ticker = {row["ticker"]: row for row in rows}
    for ticker in universe:
        row = by_ticker[ticker]
        states = "  ".join("🟢" if row[key] else "⚪" for key in ("above20", "above50", "above200"))
        output.append([
            labels.get(ticker, f"{ticker}: ETF"), f"{row['close']:.2f}", percent(row["dailyPct"]), None,
            percent(row["oneMonthPct"]), None, None, None, None, None,
            percent(row["distanceFrom52wHighPct"]), None, f"{row['rsi14']:.2f}", states, status(row),
        ])
    return output


def cell_text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value)


def workbook_values(workbook, name):
    return [
        [cell_text(value) for value in row]
        for row in workbook[name].iter_rows(values_only=True)
        if any(value is not None for value in row)
    ]


def main():
    current = json.loads((ROOT / "data" / "2026-08-20-google-sheet.json").read_text(encoding="utf-8"))
    technical = json.loads((ROOT / "data" / "2026-08-19-longbridge-adjusted.json").read_text(encoding="utf-8"))
    rows = technical["rows"]

    labels = {}
    for section in (current["sectorDashboard"], current["thematicSectors"]):
        for row in section["values"]:
            ticker = ticker_from_label(row[0] if row else None)
            if ticker:
                labels[ticker] = row[0]

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    expected_raw = workbook_values(workbook, "Weekly Expected Move")
    expected = expected_raw[:3] + [row for row in expected_raw[3:] if row and row[0] and len(row) > 8 and row[8]]
    macro = workbook_values(workbook, "Macro")
    data_qa = workbook_values(workbook, "Data QA")

    market_breadth = current["marketBreadth"]["values"][:2] + [
        row for row in current["marketBreadth"]["values"][2:] if row and row[0] and row[0] <= "2026-08-18"
    ]
    stockbee = current["stockbee"]["values"][:1] + [
        row for row in current["stockbee"]["values"][1:] if row and row[0] and row[0] != "8/19/2026"
    ]

    output = {
        "generatedAt": "2026-08-19T13:29:00Z",
        "asOf": "2026-08-19",
        "marketWatchSpreadsheetId": current["marketWatchSpreadsheetId"],
        "stockbeeSpreadsheetId": current["stockbeeSpreadsheetId"],
        "sectorDashboard": {
            **{key: value for key, value in current["sectorDashboard"].items() if key != "values"},
            "values": sheet_rows(rows, labels, technical["universes"]["sectors"]),
        },
        "thematicSectors": {
            **{key: value for key, value in current["thematicSectors"].items() if key != "values"},
            "values": sheet_rows(rows, labels, technical["universes"]["themes"]),
        },
        "macro": {**{key: value for key, value in current["macro"].items() if key != "values"}, "values": macro},
        "marketBreadth": {**{key: value for key, value in current["marketBreadth"].items() if key != "values"}, "values": market_breadth},
        "weeklyExpectedMove": {**{key: value for key, value in current["weeklyExpectedMove"].items() if key != "values"}, "values": expected},
        "dataQa": {**{key: value for key, value in current["dataQa"].items() if key != "values"}, "values": data_qa},
        "stockbee": {**{key: value for key, value in current["stockbee"].items() if key != "values"}, "values": stockbee},
    }
    destination = ROOT / "data" / "2026-08-19-google-sheet.json"
    destination.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "file": str(destination),
        "sectorRows": len(output["sectorDashboard"]["values"]) - 1,
        "thematicRows": len(output["thematicSectors"]["values"]) - 1,
        "breadthLatest": output["marketBreadth"]["values"][2],
        "stockbeeLatest": output["stockbee"]["values"][1],
        "expectedRows": len(expected) - 3,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
